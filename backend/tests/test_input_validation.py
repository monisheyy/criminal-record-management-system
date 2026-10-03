"""Input validation, pagination contracts and export safety."""
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest


def _bearer(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.parametrize("field,value", [
    ("photo_url", "javascript:alert(1)"),
    ("photo_url", "data:text/html;base64,PHNjcmlwdD4="),
    ("first_name", "<script>"),
    ("first_name", "Robert1"),
    ("crime_type", "Jaywalking"),
    ("gender", "Toaster"),
    ("prior_convictions", -1),
    ("prior_convictions", 1000),
    ("phone", "call-me-maybe"),
    ("email", "not-an-email"),
    ("date_of_birth", (datetime.now(timezone.utc) + timedelta(days=2)).isoformat()),
    ("gang_id", 999999),
])
def test_invalid_criminal_input_is_rejected(client, admin_token, field, value):
    payload = {"first_name": "Valid", "last_name": "Person", field: value}
    r = client.post("/api/criminals", json=payload, headers=_bearer(admin_token))
    assert r.status_code == 422, f"{field}={value!r} accepted: {r.text}"


def test_international_names_and_normalised_vocabularies_are_accepted(client, admin_token):
    r = client.post("/api/criminals", json={"first_name": "José", "last_name": f"O'Brien-Núñez",
                                            "crime_type": "drug trafficking", "gender": "female"},
                    headers=_bearer(admin_token))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["crime_type"] == "Drug Trafficking" and body["gender"] == "Female"
    assert body["crime_category"] == "Narcotics", "category is derived from the crime type"


def test_case_priority_from_legacy_ui_values_is_normalised(client, admin_token):
    r = client.post("/api/cases", json={"title": "Priority Normalisation", "priority": "Medium", "crime_type": "Other"},
                    headers=_bearer(admin_token))
    assert r.status_code == 200, r.text
    assert r.json()["priority"] == "normal"


def test_fir_date_cannot_precede_incident(client, admin_token):
    r = client.post("/api/cases", json={"title": "Date Order", "incident_date": "2026-03-10T10:00:00Z",
                                        "fir_date": "2026-03-01T10:00:00Z"}, headers=_bearer(admin_token))
    assert r.status_code == 422


def test_duplicate_fir_number_conflicts(client, admin_token):
    fir = f"FIR/QA/{uuid.uuid4().hex[:6]}"
    assert client.post("/api/cases", json={"title": "First FIR", "fir_number": fir}, headers=_bearer(admin_token)).status_code == 200
    assert client.post("/api/cases", json={"title": "Second FIR", "fir_number": fir}, headers=_bearer(admin_token)).status_code == 409


def test_strong_duplicate_blocks_creation_until_acknowledged(client, admin_token):
    person = {"first_name": "Twin", "last_name": "Recordholder", "date_of_birth": "1990-05-01T00:00:00Z"}
    assert client.post("/api/criminals", json=person, headers=_bearer(admin_token)).status_code == 200
    blocked = client.post("/api/criminals", json=person, headers=_bearer(admin_token))
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["duplicates"][0]["match_score"] >= 90
    ok = client.post("/api/criminals", json={**person, "acknowledge_possible_duplicate": True}, headers=_bearer(admin_token))
    assert ok.status_code == 200


def test_list_endpoints_expose_total_and_enforce_limits(client, admin_token):
    r = client.get("/api/criminals", params={"limit": 2}, headers=_bearer(admin_token))
    assert r.status_code == 200 and len(r.json()) <= 2
    assert int(r.headers["X-Total-Count"]) >= len(r.json())
    assert client.get("/api/criminals", params={"limit": 0}, headers=_bearer(admin_token)).status_code == 422
    assert client.get("/api/criminals", params={"limit": 10_000}, headers=_bearer(admin_token)).status_code == 422
    assert client.get("/api/cases", params={"sort": "hashed_password"}, headers=_bearer(admin_token)).status_code == 422
    ordered = client.get("/api/criminals", params={"sort": "last_name", "limit": 50}, headers=_bearer(admin_token)).json()
    names = [c["last_name"] for c in ordered]
    assert names == sorted(names)


def test_search_wildcards_are_literal(client, admin_token):
    r = client.get("/api/criminals", params={"search": "%"}, headers=_bearer(admin_token))
    assert r.status_code == 200
    assert r.json() == [], "a bare '%' must not match every record"


def test_excel_exports_neutralise_formula_injection():
    from openpyxl import load_workbook
    from app.utils.reports import generate_criminal_excel

    content = generate_criminal_excel({"crn": "CRN1", "name": "=HYPERLINK(\"http://evil\",\"x\")",
                                       "crime_category": "+SUM(1,1)", "cases": [{"case_number": "@cmd"}]})
    wb = load_workbook(io.BytesIO(content))
    values = [cell.value for ws in wb.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert not any(v.startswith(("=", "+", "@")) for v in values), values

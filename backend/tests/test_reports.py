import io
from openpyxl import load_workbook
from pypdf import PdfReader

from app.utils.reports import (
    generate_criminal_report,
    generate_criminal_excel,
    generate_case_report,
    generate_case_excel,
    generate_analytics_excel,
    generate_analytics_report,
)


def _criminal():
    return {
        "crn": "CRN12345678",
        "first_name": "Test",
        "last_name": "Criminal",
        "crime_category": "Theft",
        "prior_convictions": 2,
        "gang_name": "Test Gang",
        "risk_score": 82,
        "prediction": "Burglary",
        "confidence": 91.5,
        "cases": [{"case_number": "CASE001", "status": "open", "crime_type": "Burglary", "role": "Suspect"}],
    }


def _case():
    return {
        "case_number": "CASE001",
        "fir_number": "FIR001",
        "fir_date": "2026-09-19",
        "status": "open",
        "officer_name": "Officer Test",
        "criminals": [{"criminal": {"crn": "CRN12345678", "first_name": "Test", "last_name": "Criminal"}, "role": "Suspect"}],
        "victims": [{"name": "Victim Test", "age": 31, "gender": "F", "status": "alive", "injury_description": "None"}],
        "evidence": [{"evidence_number": "E001", "type": "Document", "description": "Test evidence", "collected_by": "Officer Test", "status": "collected"}],
    }


def test_criminal_pdf_is_valid():
    content = generate_criminal_report(_criminal(), [])
    assert content.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(content)).pages) >= 1


def test_criminal_excel_contains_required_values():
    content = generate_criminal_excel(_criminal())
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb["Criminal Profile"]
    values = list(ws.values)
    assert values[0][:4] == ("CRN", "Name", "Crime Category", "Prior Convictions")
    assert values[1][0] == "CRN12345678"
    assert values[1][2] == "Theft"
    assert values[1][5] == 82
    assert "Cases" in wb.sheetnames


def test_case_pdf_is_valid_and_contains_case_sections():
    content = generate_case_report(_case())
    assert content.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(content)).pages) >= 1


def test_case_excel_contains_required_sheets_and_values():
    content = generate_case_excel(_case())
    wb = load_workbook(io.BytesIO(content), read_only=True)
    assert {"Case Summary", "Criminals", "Victims", "Evidence"}.issubset(wb.sheetnames)
    assert list(wb["Case Summary"].values)[0][0] == "Case Number"
    assert list(wb["Case Summary"].values)[1][0] == "CASE001"
    assert list(wb["Victims"].values)[0][0] == "Name"
    assert list(wb["Victims"].values)[1][0] == "Victim Test"
    assert list(wb["Evidence"].values)[0][0] == "Evidence Number"
    assert list(wb["Evidence"].values)[1][0] == "E001"


def test_analytics_exports_are_valid():
    analytics = {
        "total_criminals": 10,
        "total_cases": 12,
        "open_cases": 5,
        "high_risk_criminals": 2,
        "pending_reviews": 3,
        "unread_alerts": 4,
        "prediction_accuracy": 87.5,
        "cases_by_status": {"open": 5, "closed": 7},
        "crimes_by_type": {"Theft": 4, "Fraud": 6},
        "monthly_cases": [{"month": "Sep 2026", "cases": 12}],
        "officer_workload": [{"officer": "Officer Test", "cases": 3}],
    }
    xlsx = generate_analytics_excel(analytics)
    wb = load_workbook(io.BytesIO(xlsx), read_only=True)
    assert "KPI Summary" in wb.sheetnames
    assert list(wb["KPI Summary"].values)[0] == ("Metric", "Value")
    assert list(wb["KPI Summary"].values)[1] == ("Total Criminals", 10)

    pdf = generate_analytics_report(analytics)
    assert pdf.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(pdf)).pages) >= 1

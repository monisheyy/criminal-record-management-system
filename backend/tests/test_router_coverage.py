"""Behaviour of the gang register, user administration, settings, notifications
and AI review endpoints that the other suites exercise only lightly."""
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from starlette.websockets import WebSocketDisconnect

from app import models
from app.security import get_password_hash
from app.utils.audit import create_notification

STRONG = "Tr1cky-Passphrase-2026"


def auth(token):
    return {"Authorization": f"Bearer {token}"}


def unique(prefix):
    return f"{prefix}{uuid.uuid4().hex[:8]}"


def unique_name(prefix):
    """Person names allow letters only, so encode the uniqueness as letters."""
    return prefix + "".join(chr(ord("a") + int(c, 16)) for c in uuid.uuid4().hex[:8])


def me(client, token):
    return client.get("/api/auth/me", headers=auth(token)).json()


def make_user(db, role=models.UserRole.investigating_officer, **extra):
    name = unique("u")
    user = models.User(username=name, email=f"{name}@test.com", full_name=f"User {name}",
                       hashed_password=get_password_hash(STRONG), role=role, **extra)
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def last_audit(db, action, resource_id=None):
    q = db.query(models.AuditLog).filter(models.AuditLog.action == action)
    if resource_id is not None:
        q = q.filter(models.AuditLog.resource_id == resource_id)
    return q.order_by(models.AuditLog.id.desc()).first()


# ── Gang register ────────────────────────────────────────────────────────────
def test_gang_lifecycle_is_admin_only_and_audited(client, admin_token, clerk_token, officer_token, db):
    name = unique("Gang ")
    payload = {"name": name, "territory": "Pune", "threat_level": "high", "member_count": 12}
    assert client.post("/api/gangs", json=payload, headers=auth(officer_token)).status_code == 403
    r = client.post("/api/gangs", json=payload, headers=auth(admin_token))
    assert r.status_code == 200, r.text
    gang = r.json()
    assert gang["threat_level"] == "high"
    assert last_audit(db, "GANG_CREATED", gang["id"]) is not None

    assert client.post("/api/gangs", json={**payload, "name": name.upper()}, headers=auth(admin_token)).status_code == 409
    assert any(g["id"] == gang["id"] for g in client.get("/api/gangs", headers=auth(clerk_token)).json())

    r = client.put(f"/api/gangs/{gang['id']}", json={"threat_level": "critical", "member_count": 15},
                   headers=auth(admin_token))
    assert r.status_code == 200 and r.json()["threat_level"] == "critical"
    log = last_audit(db, "GANG_UPDATED", gang["id"])
    assert log.details["before"]["threat_level"] == "high" and log.details["after"]["threat_level"] == "critical"

    other = client.post("/api/gangs", json={"name": unique("Other ")}, headers=auth(admin_token)).json()
    assert client.put(f"/api/gangs/{other['id']}", json={"name": name}, headers=auth(admin_token)).status_code == 409
    assert client.put("/api/gangs/999999", json={"alias": "x"}, headers=auth(admin_token)).status_code == 404
    assert client.delete("/api/gangs/999999", headers=auth(admin_token)).status_code == 404
    assert client.delete(f"/api/gangs/{gang['id']}", headers=auth(clerk_token)).status_code == 403

    assert client.delete(f"/api/gangs/{other['id']}", headers=auth(admin_token)).status_code == 200
    assert last_audit(db, "GANG_DELETED", other["id"]).details["name"] == other["name"]


def test_gang_with_members_or_ai_references_cannot_be_deleted(client, admin_token, db):
    with_member = client.post("/api/gangs", json={"name": unique("Members ")}, headers=auth(admin_token)).json()
    client.post("/api/criminals", json={"first_name": "Gang", "last_name": unique_name("Member"),
                                        "gang_id": with_member["id"]}, headers=auth(admin_token))
    r = client.delete(f"/api/gangs/{with_member['id']}", headers=auth(admin_token))
    assert r.status_code == 409 and "reference this gang" in r.json()["detail"]

    predicted = client.post("/api/gangs", json={"name": unique("Predicted ")}, headers=auth(admin_token)).json()
    db.add(models.AIPrediction(predicted_gang_id=predicted["id"], predicted_crime_type="Robbery"))
    db.commit()
    r = client.delete(f"/api/gangs/{predicted['id']}", headers=auth(admin_token))
    assert r.status_code == 409 and "AI predictions" in r.json()["detail"]


# ── User administration ──────────────────────────────────────────────────────
def test_create_user_rejects_duplicates_and_weak_passwords(client, admin_token, db):
    name = unique("new")
    body = {"username": name, "email": f"{name}@test.com", "full_name": "New Officer",
            "role": "investigating_officer", "password": STRONG}
    r = client.post("/api/admin/users", json=body, headers=auth(admin_token))
    assert r.status_code == 200, r.text
    assert r.json()["must_change_password"] is True, "an admin-chosen password must be replaced"
    assert client.post("/api/admin/users", json={**body, "email": f"x{name}@test.com"},
                       headers=auth(admin_token)).status_code == 409
    assert client.post("/api/admin/users", json={**body, "username": unique("y")},
                       headers=auth(admin_token)).status_code == 409
    weak = client.post("/api/admin/users", json={**body, "username": unique("w"), "email": f"w{name}@test.com",
                                                 "password": "password1"}, headers=auth(admin_token))
    assert weak.status_code == 422


def test_update_user_role_password_unlock_and_guards(client, admin_token, db):
    user = make_user(db, failed_login_attempts=5, locked_until=datetime.now(timezone.utc) + timedelta(minutes=10))
    other = make_user(db)
    admin_id = me(client, admin_token)["id"]

    assert client.put(f"/api/admin/users/{admin_id}", json={"role": "record_clerk"},
                      headers=auth(admin_token)).status_code == 409
    assert client.put(f"/api/admin/users/{admin_id}", json={"is_active": False},
                      headers=auth(admin_token)).status_code == 409
    assert client.put(f"/api/admin/users/{user.id}", json={"email": other.email},
                      headers=auth(admin_token)).status_code == 409
    assert client.put("/api/admin/users/999999", json={"department": "X"}, headers=auth(admin_token)).status_code == 404
    assert client.put(f"/api/admin/users/{user.id}", json={"password": "password1"},
                      headers=auth(admin_token)).status_code == 422

    r = client.put(f"/api/admin/users/{user.id}", json={"unlock": True, "password": "An0ther-Strong-Phrase!"},
                   headers=auth(admin_token))
    assert r.status_code == 200, r.text
    db.refresh(user)
    assert user.locked_until is None and user.failed_login_attempts == 0 and user.must_change_password
    details = last_audit(db, "USER_UPDATED", user.id).details
    assert details["password_reset_by_admin"] and details["unlocked"] and details["sessions_invalidated"]

    r = client.put(f"/api/admin/users/{user.id}", json={"role": "record_clerk"}, headers=auth(admin_token))
    assert r.status_code == 200 and r.json()["role"] == "record_clerk"
    assert last_audit(db, "ROLE_CHANGED", user.id).details["role_change"] is True


def test_last_active_admin_cannot_be_removed(client, admin_token, db, monkeypatch):
    from app.routers import admin as admin_router
    second_admin = make_user(db, role=models.UserRole.admin)
    monkeypatch.setattr(admin_router, "_active_admin_count", lambda *a, **k: 0)
    assert client.put(f"/api/admin/users/{second_admin.id}", json={"is_active": False},
                      headers=auth(admin_token)).status_code == 409
    assert client.delete(f"/api/admin/users/{second_admin.id}", headers=auth(admin_token)).status_code == 409


def test_delete_user_rules(client, admin_token, db):
    admin_id = me(client, admin_token)["id"]
    assert client.delete(f"/api/admin/users/{admin_id}", headers=auth(admin_token)).status_code == 400
    assert client.delete("/api/admin/users/999999", headers=auth(admin_token)).status_code == 404

    referenced = make_user(db)
    db.add(models.Case(case_number=unique("CASE/"), title="Owned case", assigned_officer_id=referenced.id))
    db.commit()
    r = client.delete(f"/api/admin/users/{referenced.id}", headers=auth(admin_token))
    assert r.status_code == 409 and "Deactivate" in r.json()["detail"]

    unused = make_user(db)
    assert client.delete(f"/api/admin/users/{unused.id}", headers=auth(admin_token)).status_code == 200
    assert db.query(models.User).filter_by(id=unused.id).first() is None
    assert last_audit(db, "USER_DELETED", unused.id).details["username"] == unused.username


def test_officer_directory_lists_only_active_officers(client, officer_token, clerk_token, db):
    inactive = make_user(db, is_active=False)
    r = client.get("/api/users/officers", headers=auth(officer_token))
    assert r.status_code == 200
    ids = {u["id"] for u in r.json()}
    assert inactive.id not in ids
    assert client.get("/api/users/officers", headers=auth(clerk_token)).status_code == 403


# ── Audit log, metrics, settings, analytics exports ──────────────────────────
def test_audit_log_filters(client, admin_token, clerk_token):
    client.post("/api/gangs", json={"name": unique("Audit ")}, headers=auth(admin_token))
    r = client.get("/api/admin/audit-logs", params={"action": "GANG_CREATED", "resource_type": "gang",
                                                    "status": "success", "limit": 5}, headers=auth(admin_token))
    assert r.status_code == 200 and r.json()
    assert all(row["action"] == "GANG_CREATED" for row in r.json())
    assert int(r.headers["x-total-count"]) >= 1

    tomorrow = (datetime.now(timezone.utc) + timedelta(days=1)).date().isoformat()
    assert client.get("/api/admin/audit-logs", params={"date_from": tomorrow},
                      headers=auth(admin_token)).json() == []
    yesterday = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
    assert client.get("/api/admin/audit-logs", params={"date_from": yesterday, "date_to": tomorrow},
                      headers=auth(admin_token)).json()
    assert client.get("/api/admin/audit-logs", params={"date_from": "last tuesday"},
                      headers=auth(admin_token)).status_code == 422
    assert client.get("/api/admin/audit-logs", headers=auth(clerk_token)).status_code == 403


def test_metrics_are_admin_only(client, admin_token, officer_token):
    r = client.get("/api/admin/metrics", headers=auth(admin_token))
    assert r.status_code == 200 and isinstance(r.json(), dict)
    assert client.get("/api/admin/metrics", headers=auth(officer_token)).status_code == 403


def test_system_settings_create_update_and_validate_keys(client, admin_token, officer_token, db):
    key = f"demo.{uuid.uuid4().hex[:6]}"
    r = client.put(f"/api/admin/settings/{key}", json={"value": "on", "description": "Demo flag"},
                   headers=auth(admin_token))
    assert r.status_code == 200 and r.json()["value"] == "on"
    r = client.put(f"/api/admin/settings/{key}", json={"value": "off"}, headers=auth(admin_token))
    assert r.json()["value"] == "off" and r.json()["description"] == "Demo flag"
    assert last_audit(db, "SETTING_UPDATED").details == {"key": key, "previous_value": "on", "value": "off"}
    assert any(s["key"] == key for s in client.get("/api/admin/settings", headers=auth(admin_token)).json())
    assert client.put("/api/admin/settings/bad key!", json={"value": "x"}, headers=auth(admin_token)).status_code == 422
    assert client.put(f"/api/admin/settings/{key}", json={"value": "x"}, headers=auth(officer_token)).status_code == 403


def test_dashboard_exports(client, clerk_token, db):
    r = client.get("/api/admin/dashboard/report/excel", headers=auth(clerk_token))
    assert r.status_code == 200
    assert load_workbook(io.BytesIO(r.content), read_only=True).sheetnames
    r = client.get("/api/admin/dashboard/report/pdf", headers=auth(clerk_token))
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert last_audit(db, "ANALYTICS_REPORT_PDF_GENERATED") is not None


# ── Notifications ────────────────────────────────────────────────────────────
def test_notifications_are_targeted_and_read_per_user(client, officer_token, clerk_token, db):
    officer_id = me(client, officer_token)["id"]
    personal = create_notification(db, title=unique("Personal "), message="For the officer only", target_user_id=officer_id)
    for_clerks = create_notification(db, title=unique("Clerks "), message="For clerks", target_role="record_clerk")

    officer_ids = {n["id"] for n in client.get("/api/notifications", headers=auth(officer_token)).json()}
    clerk_ids = {n["id"] for n in client.get("/api/notifications", headers=auth(clerk_token)).json()}
    assert personal.id in officer_ids and personal.id not in clerk_ids
    assert for_clerks.id in clerk_ids and for_clerks.id not in officer_ids

    # Marking someone else's notification behaves like a missing one and records nothing.
    assert client.post(f"/api/notifications/{personal.id}/read", headers=auth(clerk_token)).status_code == 200
    assert not db.query(models.NotificationRead).filter_by(notification_id=personal.id).count()

    assert client.post(f"/api/notifications/{personal.id}/read", headers=auth(officer_token)).status_code == 200
    assert client.post(f"/api/notifications/{personal.id}/read", headers=auth(officer_token)).status_code == 200
    assert db.query(models.NotificationRead).filter_by(notification_id=personal.id).count() == 1
    unread = client.get("/api/notifications", params={"unread_only": True}, headers=auth(officer_token)).json()
    assert personal.id not in {n["id"] for n in unread}


def test_notification_websocket_authenticates_and_answers_ping(client, officer_token):
    with client.websocket_connect(f"/api/notifications/ws?token={officer_token}") as ws:
        assert ws.receive_json() == {"event": "notification.connected"}
        ws.send_text("ping")
        assert ws.receive_json() == {"event": "notification.pong"}


@pytest.mark.parametrize("path, headers", [
    ("/api/notifications/ws", {}),
    ("/api/notifications/ws?token=not-a-jwt", {}),
])
def test_notification_websocket_rejects_missing_or_bad_tokens(client, path, headers):
    with pytest.raises(WebSocketDisconnect):
        with client.websocket_connect(path, headers=headers) as ws:
            ws.receive_json()


def test_notification_websocket_rejects_foreign_origin(client, officer_token):
    with pytest.raises(WebSocketDisconnect):
        with client.websocket_connect(f"/api/notifications/ws?token={officer_token}",
                                      headers={"origin": "https://evil.example"}) as ws:
            ws.receive_json()


def test_notification_websocket_rejects_user_who_must_change_password(client, db):
    user = make_user(db)
    token = client.post("/api/auth/login", data={"username": user.username, "password": STRONG}).json()["access_token"]
    client.cookies.clear()
    user.must_change_password = True
    db.commit()
    with pytest.raises(WebSocketDisconnect):
        with client.websocket_connect(f"/api/notifications/ws?token={token}") as ws:
            ws.receive_json()


# ── AI predictions: scoping, filters and error paths ─────────────────────────
@pytest.fixture
def officer_case(client, admin_token, officer_token):
    """A case assigned to the test officer with one linked offender."""
    criminal = client.post("/api/criminals", json={"first_name": "Scoped", "last_name": unique_name("Subject"),
                                                   "crime_type": "Fraud", "prior_convictions": 1},
                           headers=auth(admin_token)).json()
    case = client.post("/api/cases", json={"title": unique("AI scope "), "crime_type": "Fraud",
                                           "location": "Karol Bagh, New Delhi"}, headers=auth(admin_token)).json()
    client.post(f"/api/cases/{case['id']}/criminals", json={"criminal_id": criminal["id"]}, headers=auth(admin_token))
    client.post(f"/api/cases/{case['id']}/assign", json={"officer_id": me(client, officer_token)["id"]},
                headers=auth(admin_token))
    return case, criminal


def test_officer_can_assess_and_review_only_assigned_subjects(client, admin_token, officer_token, officer_case):
    case, criminal = officer_case
    stranger = client.post("/api/criminals", json={"first_name": "Not", "last_name": unique_name("Mine")},
                           headers=auth(admin_token)).json()
    assert client.post("/api/ai/predict", json={"criminal_id": stranger["id"]},
                       headers=auth(officer_token)).status_code == 403

    r = client.post("/api/ai/predict", json={"criminal_id": criminal["id"], "case_id": case["id"]},
                    headers=auth(officer_token))
    assert r.status_code == 200, r.text
    pred = r.json()
    assert client.get(f"/api/ai/predictions/{pred['id']}", headers=auth(officer_token)).status_code == 200
    listed = client.get("/api/ai/predictions", params={"case_id": case["id"]}, headers=auth(officer_token)).json()
    assert [p["id"] for p in listed] == [pred["id"]]

    review = {"status": "overridden", "remarks": "Seized ledgers point to money laundering.",
              "override_crime_type": "Money Laundering"}
    r = client.post(f"/api/ai/predictions/{pred['id']}/review", json=review, headers=auth(officer_token))
    assert r.status_code == 200 and r.json()["override_crime_type"] == "Money Laundering"
    for params in ({"review_status": "overridden"}, {"risk_level": pred["risk_level"]}, {"criminal_id": criminal["id"]}):
        assert pred["id"] in {p["id"] for p in client.get("/api/ai/predictions", params=params,
                                                          headers=auth(officer_token)).json()}

    stranger_pred = client.post("/api/ai/predict", json={"criminal_id": stranger["id"]}, headers=auth(admin_token)).json()
    assert client.get(f"/api/ai/predictions/{stranger_pred['id']}", headers=auth(officer_token)).status_code == 403
    assert client.post(f"/api/ai/predictions/{stranger_pred['id']}/review", json=review,
                       headers=auth(officer_token)).status_code == 403
    assert stranger_pred["id"] not in {p["id"] for p in client.get("/api/ai/predictions", headers=auth(officer_token)).json()}


def test_prediction_request_errors(client, admin_token, officer_token, officer_case):
    case, criminal = officer_case
    unlinked = client.post("/api/criminals", json={"first_name": "Unlinked", "last_name": unique_name("Person")},
                           headers=auth(admin_token)).json()
    assert client.post("/api/ai/predict", json={"criminal_id": 999999}, headers=auth(admin_token)).status_code == 404
    assert client.post("/api/ai/predict", json={"case_id": 999999}, headers=auth(admin_token)).status_code == 404
    assert client.post("/api/ai/predict", json={"criminal_id": unlinked["id"], "case_id": case["id"]},
                       headers=auth(admin_token)).status_code == 422
    assert client.post("/api/ai/predict", json={}, headers=auth(admin_token)).status_code == 422
    other_case = client.post("/api/cases", json={"title": unique("Other "), "crime_type": "Fraud"},
                             headers=auth(admin_token)).json()
    client.post(f"/api/cases/{other_case['id']}/assign", json={"officer_id": make_officer_id(client, admin_token)},
                headers=auth(admin_token))
    assert client.post("/api/ai/predict", json={"case_id": other_case["id"]},
                       headers=auth(officer_token)).status_code == 403
    assert client.get("/api/ai/predictions/999999", headers=auth(admin_token)).status_code == 404
    review = {"status": "confirmed", "remarks": "Matches the recorded evidence."}
    assert client.post("/api/ai/predictions/999999/review", json=review, headers=auth(admin_token)).status_code == 404


def make_officer_id(client, admin_token):
    name = unique("off")
    r = client.post("/api/admin/users", json={"username": name, "email": f"{name}@test.com", "full_name": "Other Officer",
                                              "role": "investigating_officer", "password": STRONG},
                    headers=auth(admin_token))
    return r.json()["id"]


def test_ai_status_and_model_registry(client, admin_token, officer_token, clerk_token):
    status = client.get("/api/ai/status", headers=auth(officer_token)).json()
    assert status["mode"] in {"demo", "validated-candidate"} and "advisory_notice" in status
    assert client.get("/api/ai/status", headers=auth(clerk_token)).status_code == 403
    assert isinstance(client.get("/api/ai/models", headers=auth(officer_token)).json(), list)

    justification = "Restoring the previous model after a regression in review."
    assert client.post("/api/ai/models/999999/activate", json={"justification": justification},
                       headers=auth(admin_token)).status_code == 404
    assert client.post("/api/ai/models/999999/rollback", json={"justification": justification},
                       headers=auth(admin_token)).status_code == 404
    models_list = client.get("/api/ai/models", headers=auth(admin_token)).json()
    not_retired = next((m for m in models_list if (m.get("evaluation_metadata") or {}).get("candidate_status") != "retired"), None)
    if not_retired:
        assert client.post(f"/api/ai/models/{not_retired['id']}/rollback", json={"justification": justification},
                           headers=auth(admin_token)).status_code == 409
    assert client.post("/api/ai/models/1/rollback", json={"justification": justification},
                       headers=auth(officer_token)).status_code == 403
